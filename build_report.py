from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path(r"D:\health\澳洲饮食营养健康App可行性研究报告.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "EAF4EA"
PALE_GOLD = "FFF4D6"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "1F2937"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

def set_fonts(run, latin="Calibri", east="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def set_style_font(style, latin="Calibri", east="Microsoft YaHei", size=11, color=BLACK, bold=False):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)

normal = doc.styles["Normal"]
set_style_font(normal, size=11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    st = doc.styles[name]
    set_style_font(st, size=size, color=color, bold=True)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ["List Bullet", "List Number"]:
    st = doc.styles[name]
    set_style_font(st, size=11)
    st.paragraph_format.left_indent = Inches(0.5)
    st.paragraph_format.first_line_indent = Inches(-0.25)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.167

if "Callout" not in [s.name for s in doc.styles]:
    st = doc.styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(st, size=11, color=NAVY)
    st.paragraph_format.space_before = Pt(6)
    st.paragraph_format.space_after = Pt(10)
    st.paragraph_format.left_indent = Inches(0.18)
    st.paragraph_format.right_indent = Inches(0.18)
    st.paragraph_format.line_spacing = 1.10

def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths_dxa)))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(indent))
    tblind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths_dxa[i]))
            tcw.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[i] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)

def add_table(headers, rows, widths_dxa, header_fill=LIGHT, aligns=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade_cell(c, header_fill)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_fonts(r, size=9.5, color=NAVY, bold=True)
    set_repeat_table_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if aligns:
                p.alignment = aligns[i]
            r = p.add_run(str(val))
            set_fonts(r, size=9.2, color=BLACK)
    set_table_geometry(t, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.extend([rfonts, color, underline])
    txt = OxmlElement("w:t")
    txt.text = text
    run.extend([rpr, txt])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p

def add_number(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p

def add_callout(label, text, fill=PALE_BLUE):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    set_table_geometry(t, [9360])
    c = t.cell(0, 0)
    shade_cell(c, fill)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "  ")
    set_fonts(r, color=NAVY, bold=True)
    r = p.add_run(text)
    set_fonts(r, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_fonts(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_fonts(run, size=9, color=MUTED)

# Running furniture
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("澳洲饮食营养健康App可行性研究")
set_fonts(hr, size=9, color=MUTED)
add_page_number(sec.footer.paragraphs[0])

# Cover: editorial_cover pattern, no title border.
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
r = p.add_run("市场与产品可行性研究")
set_fonts(r, size=11, color=BLUE, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run("澳洲饮食营养健康App\n可行性研究报告")
set_fonts(r, size=28, color=NAVY, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(28)
r = p.add_run("以薄荷健康类产品为参照的市场、数据、产品与监管评估")
set_fonts(r, size=14, color=DARK_BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
r = p.add_run("版本 1.0  |  研究基准日：2026年8月28日")
set_fonts(r, size=10.5, color=MUTED, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("用途：后续产品定义、用户验证、商业评估与Go/No-Go决策")
set_fonts(r, size=9.5, color=MUTED, italic=True)
doc.add_page_break()

doc.add_heading("执行摘要", level=1)
add_callout("核心判断", "在澳洲推出饮食营养健康App具有有条件可行性。机会不在于复制另一个英文卡路里记录器，而在于结合澳洲可信食品数据、中餐及亚洲餐识别、中英双语解释和可执行的下一餐搭配建议。", PALE_GREEN)
add_table(
    ["评估维度", "评分", "判断"],
    [
        ("技术可行性", "8/10", "官方数据可用，OCR、图片和语音交互成熟"),
        ("用户需求", "7/10", "体重和饮食健康需求强，但记录摩擦明显"),
        ("差异化空间", "6.5/10", "普通记录市场拥挤，亚洲饮食和双语服务仍有空档"),
        ("商业化", "5.5/10", "强免费竞品压低单纯订阅的付费空间"),
        ("监管可控性", "7/10", "保持一般健康管理定位可降低TGA门槛"),
        ("综合判断", "6.8/10", "建议先验证细分市场，再决定是否扩大投入"),
    ],
    [2100, 1100, 6160],
    aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
)
doc.add_paragraph("推荐切入点：澳洲多元文化饮食人群，首先服务经常食用中餐或亚洲餐的华语用户，再向其他亚洲及移民饮食场景扩展。")

doc.add_heading("1. 研究范围与前提", level=1)
doc.add_paragraph("本报告把可行性定义为：在澳洲面向普通消费者推出饮食记录、营养分析和搭配建议App，并在不进入疾病诊断或治疗的前提下，验证其产品、市场、数据、商业及监管成立条件。")
doc.add_heading("1.1 分析边界", level=2)
for item in [
    "研究对象为一般健康和体重管理工具，不等同于医疗器械或临床营养治疗软件。",
    "中国大陆产品信息主要来自官方App Store页面；其中用户数、食物库规模和识别准确率属于开发者自述，需在尽调阶段独立核验。",
    "本报告是桌面研究，不替代澳洲法律、隐私、税务、医疗器械和营养专业意见。",
    "收入、获客成本和留存率尚未经过真实用户测试，因此商业可行性为条件性判断。",
]: add_bullet(item)
doc.add_heading("1.2 评估框架", level=2)
doc.add_paragraph("评估从五个方面展开：大陆产品模式、本地竞争与需求、数据与技术、监管与隐私、商业模式及验证门槛。")

doc.add_heading("2. 中国大陆产品模式及启示", level=1)
doc.add_heading("2.1 薄荷健康：从食品库到完整体重管理平台", level=2)
doc.add_paragraph("薄荷健康已形成食品查询、三餐及运动记录、体重体脂跟踪、AI计划、专家服务、社区和健康商品销售的闭环。其官方页面称食品库超过100万条，以中国食物为主，并采用免费功能、会员和服务商品组合变现。")
for item in ["核心资产是长期积累的中国食品和菜肴数据，而不只是热量计算公式。", "产品围绕记录、反馈、计划、陪伴和购买形成连续使用链路。", "中餐菜名、份量和烹饪方式的本地化降低了记录门槛。"]: add_bullet(item)
doc.add_heading("2.2 好轻：硬件数据入口与AI管理", level=2)
doc.add_paragraph("好轻以体脂秤和体重数据为入口，叠加拍照识别、AI问答、动态减重计划、饮食和运动建议。该模式说明自动采集身体数据可以提高持续使用，但自有硬件会增加库存、售后和认证负担，不宜作为澳洲初创版本的第一步。")
doc.add_heading("2.3 Keep：运动生态带动饮食", level=2)
doc.add_paragraph("Keep把饮食置于吃、练、睡一体化数据体系中。其优势来自运动课程、穿戴设备和社区规模。若新产品没有运动内容生态，正面复制Keep的综合路线成本过高。")
doc.add_heading("2.4 新一代轻量产品：多模态低摩擦记录", level=2)
doc.add_paragraph("中国App Store已出现大量拍照、语音、自然语言和标签OCR记录工具，竞争重点从食品库数量转向记录速度和解释质量。但单张图片无法可靠获知用油、酱料、隐藏食材和真实重量，因此输出应被设计为可编辑估算，而非精确测量。")
add_callout("可借鉴的产品原则", "数据库本地化、多模态快速输入、可纠正估算、连续反馈和可信来源展示值得借鉴；社区、电商和自有硬件应推迟到产品市场匹配之后。")

doc.add_heading("3. 澳洲市场、竞争与需求", level=1)
doc.add_heading("3.1 主要竞品", level=2)
add_table(
    ["产品", "核心优势", "对新产品的约束", "潜在空档"],
    [
        ("Easy Diet Diary", "澳洲食品库、条码、营养素、免费无广告、专业端连接", "本地记录基础功能已被很好覆盖", "对中餐和双语文化饮食的理解可继续提升"),
        ("MyFitnessPal", "全球品牌、设备生态、宏量营养和高级记录方式", "用户规模和功能广度形成壁垒", "国家数据混杂、复杂菜肴记录仍可能繁琐"),
        ("FoodSwitch", "澳洲科研背景、条码、健康替代、免费", "超市扫描和替代推荐并非空白", "不以整日饮食和复合餐记录为核心"),
        ("Yuka", "简单评分、庞大商品库、替代建议", "健康评分交互已经成熟", "澳洲文化餐食和连续膳食建议有限"),
        ("AI记录新秀", "图片、语音、文字输入快速", "AI交互容易被复制", "可信数据、可解释性和澳洲专业审核仍可差异化"),
    ],
    [1500, 2600, 2600, 2660],
)
doc.add_heading("3.2 需求证据", level=2)
doc.add_paragraph("澳洲统计局数据显示，2022年65.8%的成年人属于超重或肥胖，其中31.7%属于肥胖。澳洲饮食App研究则发现，常用功能包括饮食记录、营养查询和条码扫描；停止使用的首要原因是耗时，其后包括难用、昂贵和用途有限。")
add_callout("市场含义", "澳洲不缺营养App，缺的是更低摩擦、更可信且更贴近日常饮食的体验。单纯增加一个卡路里计算器不能形成产品市场匹配。", PALE_GOLD)
doc.add_heading("3.3 华语及亚洲饮食切口", level=2)
doc.add_paragraph("2021年澳洲约139万人报告华人血统；约68.5万人在家使用普通话，约29.5万人使用粤语。华语用户适合作为首批验证对象，但最终市场应包括所有经常食用中餐、东南亚餐、印度餐及其他复合文化饮食的人群。")
for item in ["中西混合家庭和新移民需要双语食品名称及份量解释。", "家庭合菜、火锅、外卖和餐馆菜较难用传统数据库逐项记录。", "营养师需要看懂用户真实文化饮食，而不是只看到模糊的通用菜名。"]: add_bullet(item)

doc.add_heading("4. 数据与技术可行性", level=1)
doc.add_heading("4.1 澳洲可用数据基础", level=2)
add_table(
    ["数据层", "用途", "优势", "限制及处理"],
    [
        ("AFCD Release 3", "基础食材营养底库", "1,588种常见食品；每种至少58项核心营养数据；以实验室分析为主", "需遵守署名、相同方式共享及局限性说明"),
        ("AUSNUT 2023", "调查消费食物和复合食物映射", "更接近实际食用场景", "需核对具体许可和配方推导方式"),
        ("FSANZ品牌食品库", "澳洲包装食品", "由GS1和品牌商参与、具有验证规则", "覆盖范围和更新时效不能预设为完整"),
        ("标签OCR及用户提交", "补充新品牌和亚洲超市商品", "扩展快", "必须保存标签证据、版本、审核状态和纠错机制"),
        ("文化菜肴配方库", "中餐和亚洲复合餐", "形成差异化", "只能给区间或可编辑估算，需由专业人员审核"),
    ],
    [1650, 2000, 2700, 3010],
)
doc.add_heading("4.2 推荐的数据可信度设计", level=2)
for item in [
    "每条食品记录显示来源、适用国家、采集日期、最后审核日期和可信等级。",
    "把实验室数据、品牌标签数据、标准配方估算和用户提交明确区分。",
    "复杂菜肴先识别食物类别，再让用户快速确认份量、烹饪方式、油和酱料。",
    "保存修正反馈，用于提升检索排序和模型表现，但不能未经审核直接污染主数据库。",
    "同时支持kJ和kcal，并优先遵循澳洲标签和份量表达习惯。",
]: add_number(item)
doc.add_heading("4.3 AI边界", level=2)
doc.add_paragraph("AI适合降低输入成本、拆解餐食、解释营养数字和生成候选搭配，但不应替代营养数据库、份量确认或专业审核。建议向用户显示估算范围和不确定性，并允许一键修正。")

doc.add_heading("5. 推荐产品定位与MVP", level=1)
add_callout("建议定位", "面向澳洲多元饮食人群的双语智能饮食助手：理解本地包装食品、中餐和亚洲餐，并把营养数字转化为下一餐怎么搭配。", PALE_GREEN)
doc.add_heading("5.1 第一版必须具备", level=2)
for item in [
    "中英文食品名称、俗称和同义词搜索。",
    "澳洲包装食品条码和营养标签OCR。",
    "照片、文字和语音联合记录中餐及亚洲餐。",
    "碗、勺、片、杯和克等份量快速确认。",
    "能量、蛋白质、纤维、钠、饱和脂肪和糖等核心提示。",
    "依据澳洲饮食指南评价一天的饮食结构，并给出下一餐建议。",
    "常吃餐食复用、Apple Health或Health Connect同步、周报导出。",
    "明确显示数据来源、更新时间和估算置信度。",
]: add_bullet(item)
doc.add_heading("5.2 第一版不应建设", level=2)
for item in ["社区信息流和复杂内容审核体系。", "自营健康食品商城或自有体脂秤。", "疾病诊断、治疗或自动临床膳食方案。", "不可解释的单一AI健康分数。", "完全依赖照片得出的精确热量。"]: add_bullet(item)

doc.add_heading("6. 监管、专业与隐私", level=1)
doc.add_heading("6.1 TGA边界", level=2)
doc.add_paragraph("一般卡路里计数、健康饮食和生活方式工具通常不属于医疗器械；若软件声称诊断、预防、监测或治疗疾病，则可能成为软件医疗器械并需要纳入ARTG。对澳洲用户开放的海外软件也可能被视为在澳洲供应。")
add_table(
    ["建议使用的表达", "高风险表达"],
    [
        ("支持健康饮食；一般健康信息；帮助形成习惯；营养估算", "治疗糖尿病；逆转脂肪肝；预防疾病；自动制定治疗饮食"),
        ("对健康成年人提供一般搭配提示", "根据化验结果或疾病状态给出治疗决策"),
    ],
    [4680, 4680],
)
doc.add_heading("6.2 隐私与跨境数据", level=2)
doc.add_paragraph("饮食、体重、身体指标、疾病信息和餐食照片可能构成健康信息。若把数据发送至境外AI或云服务，澳洲机构通常需要采取合理措施确保境外接收者符合Australian Privacy Principles，并可能对境外处理行为承担责任。")
for item in ["优先选择澳洲区域存储，并记录所有数据处理地点。", "允许用户删除原始餐食照片并导出或删除账户数据。", "不把健康数据用于广告画像或未经同意的二次用途。", "采用最小化收集、明确同意、访问控制、审计日志和数据保留期限。", "上线前完成隐私影响评估，并由澳洲律师复核隐私政策和跨境安排。"]: add_bullet(item)
doc.add_heading("6.3 专业审核", level=2)
doc.add_paragraph("营养内容和搭配逻辑应由澳洲Accredited Practising Dietitian审核。涉及孕期、儿童、进食障碍风险、肾病、糖尿病、过敏和其他特殊人群时，应设置明确的转介和限制路径。")

doc.add_heading("7. 商业模式评估", level=1)
doc.add_heading("7.1 B2C订阅", level=2)
doc.add_paragraph("单纯依赖订阅存在压力，因为Easy Diet Diary和FoodSwitch提供强免费产品，MyFitnessPal则占据综合高级订阅位置。免费层可覆盖基础搜索、记录和每日概览；付费层应聚焦无限AI识别、双语周报、家庭账户、高级计划和营养师协作。")
doc.add_paragraph("建议测试每月7.99至12.99澳元或每年59至89澳元的价格区间；这仅是实验区间，不是已验证市场价格。")
doc.add_heading("7.2 B2B2C", level=2)
doc.add_paragraph("专业端可能比纯消费者订阅更具防御性，可服务APD、华人诊所、健身房、大学国际学生健康服务和企业健康项目。价值点是让专业人员快速理解中餐、家庭合菜和亚洲调料，并获得可审核的饮食周报。")
doc.add_heading("7.3 暂不建议的电商路线", level=2)
doc.add_paragraph("薄荷健康的商品和硬件闭环依赖中国供应链及规模。澳洲物流、库存和售后成本更高，过早销售自有商品还可能削弱营养建议的独立性。")

doc.add_heading("8. 12周验证计划", level=1)
doc.add_heading("阶段一：第1至3周，问题验证", level=2)
doc.add_paragraph("访谈30至50名澳洲用户，其中至少一半经常吃中餐或亚洲餐；同时访谈5至10名APD和若干诊所或健身专业人员。要求受访者实际演示最近一周的记录流程。")
doc.add_heading("阶段二：第4至8周，可用原型", level=2)
doc.add_paragraph("原型只覆盖300至500种高频基础食物、500至1,000种澳洲包装食品和100至200种常见中餐或亚洲菜，提供照片、文字、标签三种入口以及当日分析和下一餐建议。")
doc.add_page_break()
doc.add_heading("阶段三：第9至12周，付费验证", level=2)
add_table(
    ["指标", "建议Go门槛", "验证目的"],
    [
        ("常见餐食成功记录率", ">90%", "验证数据库和输入覆盖"),
        ("单餐中位记录时间", "<30秒", "验证是否真正降低摩擦"),
        ("AI结果大幅修改率", "<20%", "验证识别及份量流程"),
        ("四周留存", "20%至25%", "验证持续价值"),
        ("目标价格付费意愿", ">=5%活跃测试者", "验证消费者商业化"),
        ("专业机构试点", ">=3家", "验证B2B2C路径"),
    ],
    [2700, 2200, 4460],
)
doc.add_paragraph("上述数值是本项目建议的内部决策门槛，并非行业平均基准。正式决策还应结合获客成本、退款率、安全事件和营养师支持成本。")

doc.add_heading("9. Go/No-Go结论", level=1)
doc.add_heading("建议Go的条件", level=2)
for item in ["中餐及亚洲餐记录成功率显著优于通用竞品。", "用户能够在30秒内完成多数餐食记录并愿意持续四周。", "可信来源和可编辑估算能够建立用户信任。", "至少一个商业路径在小规模测试中出现真实付费。", "隐私、跨境数据和TGA定位经澳洲专业顾问确认可控。"]: add_bullet(item)
doc.add_heading("建议No-Go或转向的信号", level=2)
for item in ["主要用户只需要现有免费记录功能。", "高频中餐仍需大量手动修正，无法形成体验优势。", "付费意愿不足以覆盖AI推理、数据维护和专业审核成本。", "获客高度依赖持续付费广告，缺乏营养师或社区渠道。", "产品为追求差异化被迫进入未经验证的疾病治疗建议。"]: add_bullet(item)
add_callout("最终结论", "最值得做的是澳洲本地标准下、真正懂中餐和亚洲饮食的双语营养记录与搭配助手。最不值得做的是把薄荷健康翻译成英文后接入一个拍照模型。", PALE_GREEN)

doc.add_heading("10. 后续评估清单", level=1)
add_table(
    ["待验证问题", "当前状态", "所需证据", "责任人/日期"],
    [
        ("首批用户是否愿为双语和文化餐食识别付费？", "未验证", "付费落地页及价格实验", "待定"),
        ("澳洲品牌和亚洲超市商品覆盖率能否达到90%？", "未验证", "条码样本审计", "待定"),
        ("中餐份量和用油确认是否足够简洁？", "未验证", "可用性测试录像", "待定"),
        ("AFCD及其他数据许可如何影响数据库分发？", "待法律复核", "许可意见和数据架构", "待定"),
        ("B2B营养师端能否形成稳定收入？", "未验证", "3家以上付费试点", "待定"),
        ("目标用户中是否存在进食障碍或过度记录风险？", "需专项设计", "安全评估和转介策略", "待定"),
    ],
    [3300, 1400, 3000, 1660],
)
doc.add_paragraph("建议每轮用户研究后更新本节，同时维护：假设、证据、决策、负责人和复查日期。")

doc.add_heading("附录A：主要资料来源", level=1)
sources = [
    ("薄荷健康 App Store", "https://apps.apple.com/cn/app/id457856023"),
    ("好轻 App Store", "https://apps.apple.com/cn/app/id889565307"),
    ("Keep App Store", "https://apps.apple.com/cn/app/id952694580"),
    ("FSANZ Food and nutrient databases", "https://www.foodstandards.gov.au/science-data/food-nutrient-databases"),
    ("AFCD data files", "https://www.foodstandards.gov.au/science-data/food-nutrient-databases/afcd/data-files"),
    ("FSANZ Data User Licence Agreement", "https://www.foodstandards.gov.au/science-data/monitoringnutrients/afcd/datauserlicenceagreement"),
    ("Australian Branded Food Database", "https://www.foodstandards.gov.au/science-data/food-nutrient-databases/branded-food-database"),
    ("ABS: Waist circumference and BMI, 2022", "https://www.abs.gov.au/statistics/health/health-conditions-and-risks/waist-circumference-and-bmi/latest-release"),
    ("ABS: Cultural diversity, Census 2021", "https://www.abs.gov.au/statistics/people/people-and-communities/cultural-diversity-census/2021"),
    ("Australian Government: Eating well", "https://www.health.gov.au/topics/food-and-nutrition/about/eating-well"),
    ("NHMRC Nutrient Reference Values", "https://www.nhmrc.gov.au/about-us/publications/nutrient-reference-values-australia-and-new-zealand-including-recommended-dietary-intakes"),
    ("TGA software-based medical device guidance", "https://www.tga.gov.au/resources/guidance/understanding-how-we-regulate-software-based-medical-devices"),
    ("OAIC health information guidance", "https://www.oaic.gov.au/privacy/privacy-legislation/the-privacy-act/health-and-and-medical-research"),
    ("OAIC APP 8 cross-border disclosure", "https://www.oaic.gov.au/privacy/australian-privacy-principles/australian-privacy-principles-guidelines/chapter-8-app-8-cross-border-disclosure-of-personal-information"),
    ("Healthdirect: Dietitians", "https://www.healthdirect.gov.au/dietitians"),
    ("Australian diet-app acceptability study", "https://journals.sagepub.com/doi/10.1177/20552076221139091"),
    ("JMIR: AI image-based dietary assessment review", "https://www.jmir.org/2024/1/e51432/"),
    ("Easy Diet Diary Australia", "https://apps.apple.com/au/app/easy-diet-diary/id436104108"),
    ("MyFitnessPal Australia", "https://apps.apple.com/au/app/myfitnesspal-calorie-counter/id341232718"),
    ("FoodSwitch Australia", "https://apps.apple.com/au/app/foodswitch/id1059284559"),
    ("Yuka Australia", "https://apps.apple.com/au/app/yuka-food-cosmetic-scanner/id1092799236"),
]
for i, (label, url) in enumerate(sources, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_hyperlink(p, label, url)

doc.add_heading("附录B：版本记录", level=1)
add_table(
    ["版本", "日期", "说明"],
    [("1.0", "2026-08-28", "初始桌面研究；建立市场、数据、产品、监管、商业及验证框架")],
    [1200, 1800, 6360],
)

# Document metadata
doc.core_properties.title = "澳洲饮食营养健康App可行性研究报告"
doc.core_properties.subject = "市场、产品、数据、监管与商业可行性"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "澳洲, 营养App, 薄荷健康, 可行性, 中餐, 双语"

# Avoid table rows splitting across pages.
for table in doc.tables:
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        trpr.append(cant_split)

doc.save(OUT)
print(OUT)
