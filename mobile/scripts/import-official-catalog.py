"""Build the bundled FSANZ + USDA catalog from official download files."""
from __future__ import annotations

import json
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data-raw"
OUT = ROOT / "src" / "data" / "generated"
KJ_TO_KCAL = 4.184
UPDATED = "2026-09-04"

FSANZ_LIMITATION = (
    "There are limitations associated with food composition databases. "
    "Food composition data used in the database or databases may represent an average of the nutrient content "
    "of a particular sample of foods and ingredients, determined at a particular time. "
    "The nutrient composition of foods and ingredients can vary substantially between batches and brands because of a number of factors, "
    "including changes in season, processing practices and ingredient source, and methods of calculation. "
    "The Work is based on Australian data and Australia data may not be appropriate for use in other countries."
)

USDA_NUTRIENT_IDS = {
    "energyKcal": {1008, 2047, 2048},
    "proteinG": {1003},
    "fatG": {1004},
    "carbsG": {1005},
    "fibreG": {1079},
    "sodiumMg": {1093},
    "saturatedFatG": {1258},
    "sugarG": {2000, 1063},
}


def clean_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def to_number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value) if value == value else 0.0
    text = str(value).strip().replace(",", "")
    if not text or text in {".", "-"}:
        return 0.0
    if text.startswith("<"):
        try:
            return float(text[1:]) / 2
        except ValueError:
            return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def round_nutrient(key: str, value: float) -> float:
    if value < 0:
        value = 0.0
    if key in {"energyKcal", "sodiumMg"}:
        return round(value, 1 if key == "energyKcal" else 0)
    return round(value, 2)


def nutrients(**values: float) -> dict:
    keys = ["energyKcal", "proteinG", "carbsG", "fatG", "fibreG", "sodiumMg", "saturatedFatG", "sugarG"]
    return {key: round_nutrient(key, float(values.get(key, 0) or 0)) for key in keys}


def from_kj(kj: object) -> float:
    return to_number(kj) / KJ_TO_KCAL


def slug(text: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return value or "food"


GENERIC_TOKENS = {
    "and", "with", "from", "the", "for", "raw", "fresh", "cooked", "commercial", "all", "other",
    "stir", "fry", "fried", "homemade", "lean", "strips", "diced", "added", "without",
    "style", "type", "recipe", "recipes", "plain", "including", "using", "based",
    "piece", "serve", "small", "medium", "large", "fluid", "regular",
}


def tokens(text: str) -> set[str]:
    return {part for part in re.findall(r"[a-z]{3,}", text.lower()) if part not in GENERIC_TOKENS}


def normalize_en(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def inherited_aliases(aliases: list[str]) -> list[str]:
    kept: list[str] = []
    for alias in aliases:
        compact = alias.strip()
        if not compact:
            continue
        if re.search(r"[\u4e00-\u9fff]", compact):
            kept.append(compact)
            continue
        if " " in compact and len(compact) >= 10:
            kept.append(compact)
    return kept


def category_for(group: str, name: str) -> str:
    text = f"{group} {name}".lower()
    if "fruiting vegetable" in text or "vegetable" in text or "salad" in text or "leaf" in text or "cabbage" in text or "legume" in text and "dish" not in text:
        if any(word in text for word in ["fruit juice", "fruit drink"]):
            return "fruit"
        if "fruit" in text and "vegetable" not in text and "fruiting" not in text:
            return "fruit"
        return "vegetable"
    if any(word in text for word in ["fruit juice", "fruit,", "berry", "tropical fruit", "pome", "citrus", "stone fruit", "banana", "apple"]):
        if "vegetable" not in text:
            return "fruit"
    if any(word in text for word in ["milk", "yoghurt", "yogurt", "cheese", "cream", "dairy", "custard", "ice cream", "butter"]):
        if "peanut" not in text and "coconut cream" not in text:
            return "dairy"
    if any(word in text for word in ["bread", "rice", "pasta", "noodle", "cereal", "grain", "oat", "flour", "breakfast biscuit", "weet"]):
        return "staple"
    if "potato" in text and "chip" not in text and "crisp" not in text and "dish" not in text:
        return "staple"
    if any(word in text for word in ["beef", "lamb", "pork", "chicken", "fish", "seafood", "egg", "tofu", "meat", "prawn", "tuna", "turkey", "bacon"]):
        if not any(word in text for word in ["dish", "pie", "burger", "sandwich", "pizza", "curry", "stir", "soup", "casserole", "stew", "pasta/noodle"]):
            return "protein"
    if any(word in text for word in ["biscuit", "cake", "chocolate", "confection", "snack", "chip", "crisp", "drink", "beverage", "beer", "wine", "coffee", "tea", "soft drink", "juice"]):
        return "snack"
    if any(word in text for word in ["dish", "burger", "sandwich", "pizza", "pie", "curry", "stir", "soup", "casserole", "stew", "mixed", "meal", "dumpling", "sushi", "salad,"]):
        return "mixed"
    return "mixed"


def source_from_derivation(derivation: str, dataset: str, extra: str) -> dict:
    kind = (derivation or "").strip()
    if kind.lower() == "recipe":
        return {
            "type": "recipe",
            "label": extra,
            "region": "AU",
            "confidence": "estimate",
            "updatedAt": UPDATED,
            "dataset": dataset,
        }
    if kind.lower() in {"label data", "label"}:
        return {
            "type": "label",
            "label": extra,
            "region": "AU",
            "confidence": "estimate",
            "updatedAt": UPDATED,
            "dataset": dataset,
        }
    confidence = "high" if kind.lower() == "analysed" else "medium"
    return {
        "type": "official",
        "label": extra,
        "region": "AU",
        "confidence": confidence,
        "updatedAt": UPDATED,
        "dataset": dataset,
    }


def sheet_rows(path: Path, sheet: str, header_row: int = 3):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = None
    for index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if index < header_row:
            continue
        if index == header_row:
            headers = [clean_header(cell) for cell in row]
            continue
        if not row or row[0] in (None, ""):
            continue
        yield dict(zip(headers, row))
    wb.close()


def pick_measure(rows: list[dict]) -> tuple[str, float]:
    ranked: list[tuple[int, str, float]] = []
    for row in rows:
        grams = to_number(row.get("Gram amount"))
        if grams <= 0 or grams > 800:
            continue
        quantity = row.get("Quantity") or 1
        descriptors = " ".join(
            str(row.get(key) or "") for key in ["Descriptor 1", "Descriptor 2", "Descriptor 3", "Descriptor 4"]
        ).strip()
        lowered = descriptors.lower()
        if "density" in lowered:
            continue
        label = f"{quantity} {descriptors}".strip()
        score = 0
        if any(word in lowered for word in ["cup", "slice", "medium", "small", "fillet", "piece", "serve", "egg", "bowl", "tablespoon", "can"]):
            score += 3
        if 20 <= grams <= 400:
            score += 2
        if "100 g" in lowered or lowered == "g":
            score -= 1
        ranked.append((score, label[:48] or "100 g", grams))
    if not ranked:
        return "100 g", 100.0
    ranked.sort(key=lambda item: (-item[0], abs(item[2] - 100)))
    _, label, grams = ranked[0]
    return label, grams


def parse_seed_names() -> list[tuple[str, str, list[str]]]:
    text = (ROOT / "src" / "data" / "seed-foods.ts").read_text(encoding="utf-8")
    found: list[tuple[str, str, list[str]]] = []
    for match in re.finditer(
        r"item\('([^']+)', '([^']+)', '([^']+)', \[([^\]]*)\]",
        text,
    ):
        aliases = re.findall(r"'([^']+)'", match.group(4))
        found.append((match.group(2), match.group(3), aliases))
    return found


def chinese_for(name_en: str, seeds: list[tuple[str, str, list[str]]]) -> tuple[str, list[str]]:
    official = normalize_en(name_en)
    official_tokens = tokens(name_en)

    for name_zh, name_en_seed, aliases in seeds:
        if normalize_en(name_en_seed) == official:
            return name_zh, inherited_aliases(aliases)

    for name_zh, name_en_seed, aliases in seeds:
        seed = normalize_en(name_en_seed)
        if len(seed) < 12:
            continue
        if f" {seed} " in f" {official} ":
            return name_zh, inherited_aliases(aliases)

    best: tuple[int, str, list[str]] | None = None
    for name_zh, name_en_seed, aliases in seeds:
        seed_tokens = tokens(name_en_seed)
        if len(seed_tokens) < 2 or not seed_tokens.issubset(official_tokens):
            continue
        score = len(seed_tokens)
        if best is None or score > best[0]:
            best = (score, name_zh, inherited_aliases(aliases))
    if best:
        return best[1], best[2]
    return name_en, []


def header_get(row: dict, *candidates: str, energy: bool = False):
    for key, value in row.items():
        compact = re.sub(r"\s+", " ", key.lower())
        if not energy and compact.startswith("energy"):
            continue
        if any(candidate in compact for candidate in candidates):
            return value
    return None


def build_ausnut(seeds: list[tuple[str, str, list[str]]]) -> list[dict]:
    measures: dict[str, list[dict]] = defaultdict(list)
    for row in sheet_rows(RAW / "ausnut-food-measures.xlsx", "AUSNUT 2023"):
        survey_id = str(row.get("Survey ID") or "").strip()
        if survey_id:
            measures[survey_id].append(row)

    foods = []
    for row in sheet_rows(RAW / "ausnut-nutrient-profiles.xlsx", "Food nutrient profiles"):
        survey_id = str(row.get("Survey ID") or "").strip()
        public_key = str(row.get("Public food key") or "").strip()
        name = str(row.get("Food name") or "").strip()
        if not survey_id or not name:
            continue
        serving_label, serving_grams = pick_measure(measures.get(survey_id, []))
        name_zh, aliases = chinese_for(name, seeds)
        derivation = str(row.get("Derivation") or "")
        foods.append({
            "id": f"ausnut-{survey_id}",
            "nameZh": name_zh,
            "nameEn": name,
            "aliases": list(dict.fromkeys([*aliases, public_key, survey_id])),
            "category": "mixed",
            "servingLabel": serving_label,
            "servingGrams": serving_grams,
            "nutrientsPer100g": nutrients(
                energyKcal=from_kj(header_get(row, "energy with dietary fibre", energy=True)),
                proteinG=to_number(header_get(row, "protein (g)")),
                carbsG=to_number(header_get(row, "available carbohydrate, without sugar alcohols")),
                fatG=to_number(header_get(row, "total fat")),
                fibreG=to_number(header_get(row, "dietary fibre")),
                sodiumMg=to_number(header_get(row, "sodium (na)")),
                saturatedFatG=to_number(header_get(row, "total saturated fat")),
                sugarG=to_number(header_get(row, "total sugars")),
            ),
            "source": source_from_derivation(
                derivation,
                "fsanz-ausnut",
                f"FSANZ AUSNUT 2023（{derivation or 'published'}）。澳洲调查食物成分，按 Data User Licence 署名使用，不是医疗诊断数据。{public_key}",
            ),
            "tags": ["imported", "fsanz", "ausnut"],
            "_publicKey": public_key,
            "_group": "",
        })

    details = {str(row.get("Survey ID") or "").strip(): row for row in sheet_rows(RAW / "ausnut-food-details.xlsx", "Food details")}
    for food in foods:
        survey_id = food["id"].removeprefix("ausnut-")
        detail = details.get(survey_id) or {}
        group = str(detail.get("Food group name") or "")
        food["category"] = category_for(group, food["nameEn"])
        food["_group"] = group
        if group:
            food["tags"].append(slug(group)[:40])
        food["source"]["externalId"] = survey_id
        if detail.get("Food description"):
            food["aliases"].append(str(detail["Food description"])[:80])
        del food["_group"]
    return foods


def build_afcd(used_keys: set[str], seeds: list[tuple[str, str, list[str]]]) -> list[dict]:
    profiles = {str(row.get("Public Food Key") or "").strip(): row for row in sheet_rows(RAW / "afcd-nutrient-profiles.xlsx", "All solids & liquids per 100 g")}
    foods = []
    for row in sheet_rows(RAW / "afcd-food-details.xlsx", "Food details"):
        key = str(row.get("Public Food Key") or "").strip()
        name = str(row.get("Food Name") or "").strip()
        if not key or not name or key in used_keys:
            continue
        profile = profiles.get(key)
        if not profile:
            continue
        name_zh, aliases = chinese_for(name, seeds)
        classification = str(row.get("Classification") or "")
        derivation = str(row.get("Derivation") or "")
        foods.append({
            "id": f"afcd-{key.lower()}",
            "nameZh": name_zh,
            "nameEn": name,
            "aliases": list(dict.fromkeys([*aliases, key, classification])),
            "category": category_for(classification, name),
            "servingLabel": "100 g",
            "servingGrams": 100,
            "nutrientsPer100g": nutrients(
                energyKcal=from_kj(header_get(profile, "energy with dietary fibre", energy=True)),
                proteinG=to_number(header_get(profile, "protein")),
                carbsG=to_number(header_get(profile, "available carbohydrate, without sugar alcohols")),
                fatG=to_number(header_get(profile, "fat, total")),
                fibreG=to_number(header_get(profile, "total dietary fibre")),
                sodiumMg=to_number(header_get(profile, "sodium (na)")),
                saturatedFatG=to_number(header_get(profile, "total saturated fatty acids, equated  (g)")),
                sugarG=to_number(header_get(profile, "total sugars")),
            ),
            "source": {
                **source_from_derivation(
                    derivation,
                    "fsanz-afcd",
                    f"FSANZ AFCD Release 3（{derivation or 'published'}）。澳洲参考分析库，按 Data User Licence 署名使用。{key}",
                ),
                "externalId": key,
            },
            "tags": ["imported", "fsanz", "afcd"],
        })
    return foods


def usda_nutrients(entry: dict) -> dict:
    collected: dict[str, float] = {}
    for item in entry.get("foodNutrients") or []:
        nutrient = item.get("nutrient") or {}
        nutrient_id = nutrient.get("id")
        amount = item.get("amount")
        if amount is None:
            continue
        for key, ids in USDA_NUTRIENT_IDS.items():
            if nutrient_id in ids and key not in collected:
                collected[key] = float(amount)
    if "energyKcal" not in collected:
        for item in entry.get("foodNutrients") or []:
            nutrient = item.get("nutrient") or {}
            if nutrient.get("unitName") == "kcal" and "energy" in str(nutrient.get("name", "")).lower():
                collected["energyKcal"] = float(item.get("amount") or 0)
                break
    return nutrients(**collected)


def usda_portion(entry: dict) -> tuple[str, float]:
    portions = entry.get("foodPortions") or []
    ranked = []
    for portion in portions:
        grams = to_number(portion.get("gramWeight"))
        if grams <= 0:
            continue
        label = str(portion.get("portionDescription") or portion.get("modifier") or "portion").strip()
        ranked.append((abs(grams - 100), label, grams))
    if not ranked:
        return "100 g", 100.0
    ranked.sort()
    _, label, grams = ranked[0]
    return (label or "100 g")[:48], grams


def build_usda() -> list[dict]:
    with zipfile.ZipFile(RAW / "usda-foundation.zip") as archive:
        name = next(item for item in archive.namelist() if item.endswith(".json"))
        with archive.open(name) as handle:
            payload = json.load(handle)
    foods = []
    for entry in payload.get("FoundationFoods") or []:
        if not isinstance(entry, dict):
            continue
        fdc_id = entry.get("fdcId")
        name = str(entry.get("description") or "").strip()
        if not fdc_id or not name:
            continue
        category_name = ((entry.get("foodCategory") or {}).get("description")) or ""
        serving_label, serving_grams = usda_portion(entry)
        foods.append({
            "id": f"usda-{fdc_id}",
            "nameZh": name,
            "nameEn": name,
            "aliases": [str(fdc_id), "USDA", "美国"],
            "category": category_for(category_name, name),
            "servingLabel": serving_label,
            "servingGrams": serving_grams,
            "nutrientsPer100g": usda_nutrients(entry),
            "source": {
                "type": "official",
                "label": "USDA FoodData Central Foundation Foods，CC0 公有领域。美国分析样品，仅作对照，不是澳洲官方值。",
                "region": "US",
                "confidence": "high",
                "updatedAt": "2026-04-30",
                "dataset": "usda-fdc",
                "externalId": str(fdc_id),
            },
            "tags": ["imported", "overseas", "usda", "reference"],
        })
    return foods


def attach_references(au_foods: list[dict], usda_foods: list[dict]) -> None:
    index = [(set(tokens(food["nameEn"])), food) for food in usda_foods]
    for food in au_foods:
        au_tokens = tokens(food["nameEn"])
        if len(au_tokens) < 2:
            continue
        scored = []
        for usda_tokens, usda in index:
            overlap = au_tokens & usda_tokens
            if len(overlap) < 2:
                continue
            score = len(overlap) / min(len(au_tokens), len(usda_tokens))
            if score >= 0.5:
                scored.append((score, len(overlap), usda))
        if not scored:
            continue
        scored.sort(key=lambda item: (-item[0], -item[1]))
        match = scored[0][2]
        food["overseasReference"] = {
            "dataset": "usda-fdc",
            "nameEn": match["nameEn"],
            "externalId": match["source"]["externalId"],
            "label": "USDA Foundation 对照（美国样品，CC0）",
            "nutrientsPer100g": match["nutrientsPer100g"],
        }


def strip_private(food: dict) -> dict:
    food = dict(food)
    food.pop("_publicKey", None)
    food["source"].pop("_publicKey", None)
    food["aliases"] = [alias for alias in food["aliases"] if alias][:8]
    food["tags"] = list(dict.fromkeys(food["tags"]))
    return food


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    seeds = parse_seed_names()
    ausnut = build_ausnut(seeds)
    used_keys = {food.pop("_publicKey", "") for food in ausnut}
    used_keys.discard("")
    afcd = build_afcd(used_keys, seeds)
    usda = build_usda()
    attach_references(ausnut + afcd, usda)
    imported = [strip_private(food) for food in [*ausnut, *afcd, *usda]]
    stats = {
        "updatedAt": UPDATED,
        "ausnut": len(ausnut),
        "afcdExtra": len(afcd),
        "usdaFoundation": len(usda),
        "withUsdaReference": sum(1 for food in imported if food.get("overseasReference")),
        "totalImported": len(imported),
        "licence": {
            "fsanz": "CC BY-SA 3.0 Australia + FSANZ Data User Licence",
            "usda": "CC0 1.0",
            "limitationOfData": FSANZ_LIMITATION,
        },
    }
    (OUT / "imported-foods.ts").write_text(
        "/* eslint-disable */\n// @ts-nocheck\nexport const importedFoods = "
        + json.dumps(imported, ensure_ascii=False, separators=(",", ":"))
        + ";\n",
        encoding="utf-8",
    )
    (OUT / "import-stats.ts").write_text(
        "export const importStats = " + json.dumps(stats, ensure_ascii=False, indent=2) + " as const;\n",
        encoding="utf-8",
    )
    (OUT / "LICENCE-FSANZ.txt").write_text(
        "Australian Food Composition Database / AUSNUT 2023\n"
        "Licensed by Food Standards Australia New Zealand under a licence based on CC BY-SA 3.0 Australia.\n"
        "https://www.foodstandards.gov.au/science-data/monitoringnutrients/afcd/datauserlicenceagreement\n\n"
        f"{FSANZ_LIMITATION}\n\n"
        "USDA FoodData Central Foundation Foods are CC0 1.0 public domain.\n"
        "U.S. Department of Agriculture, Agricultural Research Service. FoodData Central, 2026. fdc.nal.usda.gov.\n",
        encoding="utf-8",
    )
    print(json.dumps(stats, indent=2))


if __name__ == "__main__":
    main()
